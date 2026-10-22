from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "documentacion" / "originales"
OUTPUT_DIR = ROOT / "documentacion" / "texto"
EXTENSIONES = {".pdf", ".docx"}


def nombre_salida(path: Path) -> Path:
    """Devuelve la ruta Markdown correspondiente al documento original."""
    return OUTPUT_DIR / f"{path.stem}.md"


def convertir_pdf(path: Path) -> str:
    """Extrae el texto de cada página de un PDF y lo separa por encabezados."""
    reader = PdfReader(str(path))
    partes = []

    for numero, page in enumerate(reader.pages, start=1):
        texto = (page.extract_text() or "").strip()
        partes.append(f"## Página {numero}\n\n")
        partes.append(
            texto if texto else "_[La página no contiene texto extraíble.]_"
        )
        partes.append("\n\n")

    return "".join(partes)


def escapar_celda_markdown(texto: str) -> str:
    """Evita que barras verticales rompan las tablas Markdown."""
    return texto.replace("\\", "\\\\").replace("|", "\\|").replace("\n", " ").strip()


def convertir_docx(path: Path) -> str:
    """Extrae párrafos y tablas de un documento DOCX."""
    doc = Document(str(path))
    partes = []

    for parrafo in doc.paragraphs:
        texto = parrafo.text.strip()
        if texto:
            partes.append(texto)

    for tabla_num, tabla in enumerate(doc.tables, start=1):
        partes.append(f"\n## Tabla {tabla_num}\n")
        filas = [
            [escapar_celda_markdown(celda.text) for celda in fila.cells]
            for fila in tabla.rows
        ]

        if not filas:
            continue

        numero_columnas = len(filas[0])
        encabezado = filas[0]
        partes.append("| " + " | ".join(encabezado) + " |")
        partes.append("| " + " | ".join(["---"] * numero_columnas) + " |")

        for fila in filas[1:]:
            # Mantener el mismo número de columnas en todas las filas.
            fila = (fila + [""] * numero_columnas)[:numero_columnas]
            partes.append("| " + " | ".join(fila) + " |")

    return "\n\n".join(partes)


def main() -> None:
    """Convierte todos los PDF y DOCX encontrados en la carpeta de originales."""
    if not SOURCE_DIR.exists():
        print(f"No existe la carpeta de documentos originales: {SOURCE_DIR}")
        return

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    documentos = sorted(
        path
        for path in SOURCE_DIR.rglob("*")
        if path.is_file() and path.suffix.lower() in EXTENSIONES
    )

    if not documentos:
        print("No se encontraron PDF o DOCX en documentacion/originales/.")
        return

    errores = []

    for path in documentos:
        try:
            if path.suffix.lower() == ".pdf":
                contenido = convertir_pdf(path)
            else:
                contenido = convertir_docx(path)

            titulo = path.stem.replace("_", " ")
            salida = nombre_salida(path)
            salida.write_text(
                f"# {titulo}\n\n"
                f"> Versión de texto generada automáticamente a partir de "
                f"{path.name}.\n\n"
                f"{contenido.rstrip()}\n",
                encoding="utf-8",
            )
            print(f"Convertido: {path.relative_to(ROOT)} -> {salida.relative_to(ROOT)}")
        except Exception as error:
            errores.append((path, error))
            print(f"ERROR al convertir {path.relative_to(ROOT)}: {error}")

    if errores:
        raise SystemExit(f"Falló la conversión de {len(errores)} documento(s).")

    print(f"Conversión completada: {len(documentos)} documento(s).")


if __name__ == "__main__":
    main()
